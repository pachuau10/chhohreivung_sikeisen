import os
import fitz
import docx
from django.views.generic import CreateView, DetailView, ListView, View
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import F, Q
from django.db import connection
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse_lazy
from django.http import HttpResponse, JsonResponse, Http404
from django.contrib.postgres.search import SearchVector, SearchQuery, SearchRank
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model
from .models import Note, Subject, Bookmark, Like, Download, Comment
User = get_user_model()
from .forms import NoteUploadForm, NoteSearchForm
from tests.models import UsageLog


class NoteUploadView(LoginRequiredMixin, CreateView):
    model = Note
    form_class = NoteUploadForm
    template_name = 'notes/note_upload.html'
    success_url = reverse_lazy('dashboard')

    def form_valid(self, form):
        user = self.request.user
        if not user.is_staff and user.profile.uploads_remaining <= 0:
            messages.error(self.request, f'Upload limit reached. Free plan allows {user.profile._get_limits()["uploads"]} uploads per day.')
            return self.form_invalid(form)
        form.instance.user = user
        file = form.cleaned_data['file']
        form.instance.file_type = file.name.split('.')[-1].lower()
        response = super().form_valid(form)
        self.process_document(self.object)
        user.profile.uploads_today += 1
        user.profile.save()
        UsageLog.objects.create(user=user, action='note_upload', detail={'note_id': self.object.id})
        messages.success(self.request, 'Note uploaded successfully!')
        return response

    def process_document(self, note):
        try:
            file_path = note.file.path
            text = ''
            if note.file_type == 'pdf':
                doc = fitz.open(file_path)
                note.page_count = doc.page_count
                for page in doc:
                    text += page.get_text()
                doc.close()
            elif note.file_type == 'docx':
                doc = docx.Document(file_path)
                text = '\n'.join([p.text for p in doc.paragraphs])
                note.page_count = len(doc.paragraphs)
            note.extracted_text = text[:50000]
            note.is_processed = True
            note.save(update_fields=['extracted_text', 'page_count', 'is_processed'])
        except Exception as e:
            import logging
            logger = logging.getLogger(__name__)
            logger.error(f'Document processing failed for note {note.id}: {e}')


class NoteDetailView(LoginRequiredMixin, DetailView):
    model = Note
    template_name = 'notes/note_detail.html'
    context_object_name = 'note'

    def get_object(self, queryset=None):
        note = get_object_or_404(Note, pk=self.kwargs.get('pk'))
        if note.visibility != 'public' and note.user != self.request.user:
            raise Http404
        return Note.objects.select_related('user', 'subject').prefetch_related(
            'flashcards', 'summaries', 'generated_questions'
        ).get(pk=note.pk)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        note = self.object
        if self.request.user.is_authenticated:
            context['is_bookmarked'] = Bookmark.objects.filter(user=self.request.user, note=note).exists()
            context['is_liked'] = Like.objects.filter(user=self.request.user, note=note).exists()
        context['mcqs'] = note.generated_questions.filter(question_type='mcq')
        context['short_questions'] = note.generated_questions.filter(question_type='short')
        context['long_questions'] = note.generated_questions.filter(question_type='long')
        context['flashcards'] = note.flashcards.all()
        context['summaries'] = note.summaries.all()
        if note.subject:
            context['related_notes'] = Note.objects.filter(visibility='public', subject=note.subject).exclude(pk=note.pk).select_related('user', 'subject')[:4]
        else:
            context['related_notes'] = Note.objects.filter(visibility='public').exclude(pk=note.pk).select_related('user', 'subject')[:4]
        return context


class NoteListView(ListView):
    model = Note
    template_name = 'notes/note_list.html'
    context_object_name = 'notes'
    paginate_by = 12

    def get_queryset(self):
        queryset = Note.objects.filter(visibility='public').select_related('user', 'subject')
        form = NoteSearchForm(self.request.GET)
        if form.is_valid():
            q = form.cleaned_data.get('q')
            subject_slug = form.cleaned_data.get('subject')
            if q:
                if connection.vendor == 'postgresql':
                    search_query = SearchQuery(q, config='english')
                    vector = SearchVector('title', weight='A') + SearchVector('description', weight='B') + SearchVector('extracted_text', weight='C')
                    queryset = queryset.annotate(rank=SearchRank(vector, search_query)).filter(rank__gte=0.01).order_by('-rank')
                else:
                    queryset = queryset.filter(
                        Q(title__icontains=q) | Q(description__icontains=q) | Q(extracted_text__icontains=q)
                    )
            if subject_slug:
                queryset = queryset.filter(subject__slug=subject_slug)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search_form'] = NoteSearchForm(self.request.GET)
        context['subjects'] = Subject.objects.all()
        if self.request.user.is_authenticated:
            context['followed_ids'] = list(self.request.user.profile.followed_subjects.values_list('pk', flat=True))
        else:
            context['followed_ids'] = []
        return context


class BookmarkToggleView(LoginRequiredMixin, View):
    def post(self, request, pk):
        note = get_object_or_404(Note, pk=pk)
        bookmark, created = Bookmark.objects.get_or_create(user=request.user, note=note)
        if not created:
            bookmark.delete()
            Note.objects.filter(pk=pk).update(bookmark_count=F('bookmark_count') - 1)
            return JsonResponse({'bookmarked': False, 'count': note.bookmark_count - 1})
        Note.objects.filter(pk=pk).update(bookmark_count=F('bookmark_count') + 1)
        return JsonResponse({'bookmarked': True, 'count': note.bookmark_count + 1})


class LikeToggleView(LoginRequiredMixin, View):
    def post(self, request, pk):
        deleted, _ = Like.objects.filter(user=request.user, note_id=pk).delete()
        if deleted:
            Note.objects.filter(pk=pk).update(like_count=F('like_count') - 1)
            count = Note.objects.values_list('like_count', flat=True).get(pk=pk)
            return JsonResponse({'liked': False, 'count': count})
        note = Note.objects.only('user_id', 'title').get(pk=pk)
        Like.objects.create(user=request.user, note_id=pk)
        Note.objects.filter(pk=pk).update(like_count=F('like_count') + 1)
        if note.user_id != request.user.pk:
            from core.views import create_notification
            create_notification(note.user, 'like', f'{request.user.username} liked your note', f'"{note.title}"', f'/notes/{pk}/')
        count = Note.objects.values_list('like_count', flat=True).get(pk=pk)
        return JsonResponse({'liked': True, 'count': count})


class NoteDownloadView(LoginRequiredMixin, View):
    def get(self, request, pk):
        note = get_object_or_404(Note, pk=pk)
        if not request.user.is_staff and request.user.profile.plan == 'free':
            messages.error(request, 'Downloading notes requires Premium.')
            return redirect('pricing')
        Download.objects.create(user=request.user, note=note)
        Note.objects.filter(pk=pk).update(download_count=F('download_count') + 1)
        UsageLog.objects.create(user=request.user, action='note_download', detail={'note_id': note.id})
        return redirect(note.file.url)


class NoteDeleteView(LoginRequiredMixin, View):
    def post(self, request, pk):
        note = get_object_or_404(Note, pk=pk, user=request.user)
        if note.file:
            try:
                if os.path.exists(note.file.path):
                    os.remove(note.file.path)
            except Exception:
                pass
        note.delete()
        messages.success(request, f'"{note.title}" deleted successfully.')
        return redirect('dashboard')


class CommentAddView(LoginRequiredMixin, View):
    def post(self, request, pk):
        note = get_object_or_404(Note, pk=pk, visibility='public')
        content = request.POST.get('content', '').strip()
        if not content:
            return JsonResponse({'error': 'Comment cannot be empty.'}, status=400)
        comment = Comment.objects.create(note=note, user=request.user, content=content)
        if note.user != request.user:
            from core.views import create_notification
            create_notification(note.user, 'comment', f'{request.user.username} commented on your note', f'"{content[:80]}"', f'/notes/{pk}/')
        return JsonResponse({
            'success': True,
            'id': comment.id,
            'user': comment.user.username,
            'content': comment.content,
            'created_at': comment.created_at.isoformat(),
        })
